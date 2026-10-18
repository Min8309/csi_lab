#!/usr/bin/env python3
"""Rebuild the CSI LAB presentation from inspected source and bundled data.
Run with Python 3.12 and python-pptx, pandas, Pillow, Pygments.
No application/data files are modified. Screenshot assets are retained captures.
"""
from pathlib import Path
import json, io, textwrap, tokenize, token, re, ast
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image
import pandas as pd
from pygments import lex
from pygments.lexers import PythonLexer
from pygments.token import Token

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/presentation'; AS=OUT/'assets'; OUT.mkdir(parents=True,exist_ok=True)
AUDIT=json.loads((OUT/'data_audit.json').read_text())
SHA=AUDIT['source_commit']; GH=f'https://github.com/Min8309/csi_lab/blob/{SHA}/'
prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
prs.core_properties.title='CSI LAB | 범죄 데이터를 읽는 수사 연구소'
prs.core_properties.subject='범죄율 이해와 데이터 패턴 읽기를 위한 교육 프로젝트'
prs.core_properties.author='CSI LAB';prs.core_properties.keywords='CSI LAB, 범죄 데이터, 데이터 리터러시, Streamlit, YOLO'
C={'bg':'F7F9F7','white':'FFFFFF','ink':'18333F','muted':'57707A','teal':'087F8C','blue':'2563B8','mint':'E0F3ED','ice':'E9F1FC','orange':'D77729','peach':'FFF0DF','line':'D7E2E4','code':'EFF4F6','red':'B84D42'}
FONT='Noto Sans CJK KR'; MONO='DejaVu Sans Mono'; narrative=[];sources=[]
def rgb(c):return RGBColor.from_string(C.get(c,c))
def box(s,x,y,w,h,fill='white',line=None,radius=False):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 sh._element.spPr.append(OxmlElement('a:effectLst'))
 if radius:
  try:sh.adjustments[0]=.12
  except Exception:pass
 if line:sh.line.color.rgb=rgb(line)
 else:sh.line.fill.background()
 return sh

def txt(s,x,y,w,h,text,size=20,color='ink',bold=False,font=FONT,align=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=Inches(.02);tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(str(text).split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(5)
  if align is not None:p.alignment=align
 return sh

def line(s,x1,y1,x2,y2,color='line',width=1.5):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(width);return sh

def pill(s,x,y,w,label,fill='mint',color='teal'):
 box(s,x,y,w,.38,fill,radius=True);txt(s,x+.12,y+.05,w-.24,.26,label,11,color,True)

def base(section,title,sub='',note='',srcs=()):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('bg');i=len(prs.slides)
 box(s,0,0,.16,9,'teal');txt(s,.65,.32,13,.35,section,12,'teal',True)
 txt(s,.65,.95,14.7,.72,title,31,'ink',True)
 if sub:txt(s,.68,1.78,14.4,.68,sub,16,'muted')
 line(s,.68,8.45,15.3,8.45);txt(s,.68,8.58,12,.22,'CSI LAB  /  CRIME DATA LITERACY  /  SOURCE '+SHA[:7],9,'muted');txt(s,14.4,8.54,.85,.3,f'{i:02}',12,'teal',True,align=PP_ALIGN.RIGHT)
 s.notes_slide.notes_text_frame.text=note+'\n\n근거:\n'+'\n'.join(srcs)
 narrative.append({'slide':i,'section':section,'title':title,'subtitle':sub,'notes':note,'sources':list(srcs)})
 return s

def card(s,x,y,w,h,kicker,title,body,fill='white',accent='teal'):
 box(s,x,y,w,h,fill,radius=True);box(s,x,y,.06,h,accent);txt(s,x+.25,y+.23,w-.5,.3,kicker,11,accent,True);txt(s,x+.25,y+.8,w-.5,.8,title,23,'ink',True);txt(s,x+.25,y+1.7,w-.5,h-1.9,body,18,'muted')

def bottom(s,text,fill='mint',color='teal',y=7.55):
 box(s,.7,y,14.6,.58,fill,radius=True);txt(s,.9,y+.12,14.2,.38,text,15,color,True)

def addpic(s,path,x,y,w,h):
 im=Image.open(path);iw,ih=im.size;r=min(w/iw,h/ih);nw,nh=iw*r,ih*r
 s.shapes.add_picture(str(path),Inches(x+(w-nw)/2),Inches(y+(h-nh)/2),width=Inches(nw),height=Inches(nh))

def code(section,title,sub,path,start,end,meaning,result,decision,note='',trim=True):
 ls=(ROOT/path).read_text().splitlines(); excerpt='\n'.join(ls[start-1:end]);excerpt=textwrap.dedent(excerpt)
 if trim:excerpt='\n'.join(x for x in excerpt.splitlines() if x.strip() and not x.lstrip().startswith('#'))
 ast.parse(excerpt)  # Every displayed excerpt must be complete Python.
 source=f'{path}:L{start}–L{end}';url=GH+path+f'#L{start}-L{end}'
 s=base(section,title,sub,note or f'{meaning}\n{result}\n{decision}',[url])
 box(s,.7,2.65,9.75,4.62,'code',radius=True);pill(s,.98,2.88,1.3,'PYTHON','ice','blue')
 sh=txt(s,2.45,2.94,7.6,.3,source,12,'muted');sh.text_frame.paragraphs[0].runs[0].hyperlink.address=url
 sh=s.shapes.add_textbox(Inches(1),Inches(3.48),Inches(9.2),Inches(3.6));tf=sh.text_frame;tf.clear();tf.word_wrap=False;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 nlines=len(excerpt.splitlines());maxlen=max(map(len,excerpt.splitlines()),default=1);fs=min(17, max(12, 230/nlines), max(12, 1050/maxlen))
 # Syntax highlighting is applied to editable text runs.
 p=tf.paragraphs[0];p.space_after=Pt(0);p.line_spacing=1.05
 for tok,value in lex(excerpt,PythonLexer()):
  for j,part in enumerate(value.split('\n')):
   if j:
    p=tf.add_paragraph();p.space_after=Pt(0);p.line_spacing=1.05
   if not part:continue
   r=p.add_run();r.text=part;r.font.name=MONO;r.font.size=Pt(fs)
   r.font.color.rgb=rgb('blue' if tok in Token.Keyword else 'teal' if tok in Token.Literal.String else 'orange' if tok in Token.Literal.Number else 'muted' if tok in Token.Comment else 'ink')
   r.font.bold=tok in Token.Keyword
 for y,label,body in [(2.75,'무엇을 구현했나',meaning),(4.18,'사용자에게 연결되는 결과',result),(5.62,'설계 판단 · 해석',decision)]:
  txt(s,10.85,y,4.35,.32,label,12,'teal',True);txt(s,10.85,y+.46,4.35,1.02,body,18,'ink')
 bottom(s,'실제 코드 발췌 · 공백·주석만 정리 · 상단 파일명을 클릭하면 원본으로 이동',fill='ice',color='blue')
 sources.append({'slide':len(prs.slides),'file':path,'start':start,'end':end,'excerpt':excerpt,'url':url})
 s.notes_slide.notes_text_frame.text+='\n\n발췌 코드:\n'+excerpt
 return s

def chart(s,x,y,w,h,categories,series,ytitle='발생건수',colors=('teal','blue')):
 data=CategoryChartData();data.categories=[str(v) for v in categories]
 for name,vals in series:data.add_series(name,vals)
 ch=s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS,Inches(x),Inches(y),Inches(w),Inches(h),data).chart
 ch.has_legend=len(series)>1
 if ch.has_legend:ch.legend.position=XL_LEGEND_POSITION.BOTTOM;ch.legend.font.name=FONT;ch.legend.font.size=Pt(12)
 ch.font.name=FONT;ch.font.size=Pt(12)
 ch.category_axis.tick_labels.font.size=Pt(11);ch.category_axis.tick_labels.font.name=FONT
 ch.value_axis.tick_labels.font.size=Pt(11);ch.value_axis.minimum_scale=0
 ch.value_axis.has_major_gridlines=True;ch.value_axis.major_gridlines.format.line.color.rgb=rgb('line')
 ch.value_axis.has_title=True;ch.value_axis.axis_title.text_frame.text=ytitle;ch.value_axis.axis_title.text_frame.paragraphs[0].font.size=Pt(11)
 for i,se in enumerate(ch.series):se.format.line.color.rgb=rgb(colors[i%len(colors)]);se.format.line.width=Pt(2.8);se.marker.style=XL_MARKER_STYLE.CIRCLE;se.marker.size=5
 return ch

# 01 — Opening
s=base('PROJECT DOSSIER  /  교육용 데이터 탐험','',note='CSI LAB은 가상 사건 탐색을 출발점으로 실제 범죄 통계를 해석하는 교육 프로젝트입니다. 핵심 목표는 범죄율의 의미를 올바르게 이해하고 범죄 데이터의 패턴을 읽는 능력입니다. 현재 앱은 발생건수와 검거율·증감률을 제공하며 인구 보정 범죄율은 계산하지 않습니다. 첨부 where_u_go.pptx는 구성과 설명 밀도의 참고자료로만 사용했습니다.',srcs=[GH+'README.md',GH+'app.py'])
# remove the empty header space by writing the main cover
pill(s,.75,1.22,2.5,'CASE FILE  /  CSI LAB');txt(s,.7,2.02,10,1.4,'CSI LAB',72,'ink',True)
txt(s,.78,3.58,10.8,1.35,'범죄 데이터를 읽는\n수사 연구소',38,'teal',True)
txt(s,.8,5.38,10.3,1.0,'범죄율을 올바르게 이해하고,\n범죄 데이터의 패턴을 읽는 교육 프로젝트',23,'ink')
# magnifier and evidence chart motif
for x,y,r,col in [(12.5,3.6,2.5,'mint'),(12.5,3.6,1.95,'white')]:
 sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-r/2),Inches(y-r/2),Inches(r),Inches(r));sh.fill.solid();sh.fill.fore_color.rgb=rgb(col);sh.line.color.rgb=rgb('teal');sh.line.width=Pt(2)
line(s,13.3,4.45,14.7,5.85,'teal',17)
for a,b,c,d in [(11.8,3.9,12.15,3.45),(12.15,3.45,12.5,3.65),(12.5,3.65,13.1,3.05)]:line(s,a,b,c,d,'blue',4)
pill(s,11.1,6.7,3.65,'OBSERVE  →  READ  →  EXPLAIN','ice','blue')
bottom(s,'가상 사건으로 몰입하고, 실제 통계로 해석을 배웁니다.')
narrative[-1]['title']='CSI LAB | 범죄 데이터를 읽는 수사 연구소'

s=base('01  PURPOSE','범죄 데이터를 이해하는 힘을 기릅니다','게임의 목적은 사건의 정답보다, 근거를 읽고 설명하는 경험입니다.',note='교육 목표는 네 가지입니다. 숫자의 단위를 구분하고, 동일 조건에서 시간 변화를 관찰하고, 증거와 해석을 분리하고, 일반적인 예방 행동으로 연결합니다. 이는 교육 설계 목표이며 학습 향상 효과는 아직 측정하지 않았습니다.',srcs=[GH+'app.py'])
for x,k,t,b,f in [( .7,'01 · READ','숫자의 의미','발생건수·검거율·증감률을\n서로 다른 지표로 읽기','mint'),(5.65,'02 · COMPARE','패턴의 맥락','지역·범죄·기간을 맞춰\n증가와 감소를 설명하기','ice'),(10.6,'03 · EXPLAIN','근거 있는 판단','관찰과 추정을 구분하고\n예방 행동으로 연결하기','peach')]:card(s,x,2.8,4.7,4.05,k,t,b,f)
bottom(s,'학습 목표: 범죄율 이해 + 데이터 패턴 읽기 + 근거를 설명하는 태도')

s=base('01  LEARNING PROBLEM','숫자를 많이 보여줘도, 이해가 생기지는 않습니다','CSI LAB은 자주 생기는 오해를 탐색 질문으로 바꿉니다.',note='발생건수가 높은 지역은 인구와 유동인구가 클 수 있습니다. 단기 증가와 장기 추세는 다를 수 있고 객체 탐지는 범죄 판단이 아닙니다. 현재 앱에 있는 경고와 시각화를 교육 질문으로 재구성했습니다.',srcs=[GH+'app.py#L2370-L2380',GH+'app.py#L3040'])
rows=[('건수가 크면 더 위험할까?','분모와 비교 조건을 확인한다.','COUNT ≠ RATE'),('올해 증가하면 장기 추세도 증가일까?','짧은 변화와 긴 흐름을 나눠 본다.','SHORT / LONG'),('가방을 탐지하면 범인을 찾은 걸까?','시각 관찰과 사건 판단을 구분한다.','OBSERVATION')]
for j,(q,a,k) in enumerate(rows):
 y=2.8+j*1.45;box(s,.7,y,14.6,1.2,'white',radius=True);pill(s,.96,y+.32,2.9,k);txt(s,4.1,y+.24,6.1,.65,q,22,'ink',True);txt(s,10.35,y+.3,4.65,.7,a,17,'muted')

s=base('01  LEARNING JOURNEY','가상 사건에서 실제 데이터로 이동합니다','CASE → EVIDENCE → AI → DATA → PREVENTION',note='CASE SOLVED 이후 REAL DATA UNLOCKED와 Safety LAB이 연결됩니다. 단, 실제 데이터는 홈과 수사 중에도 열람할 수 있으므로 데이터가 오직 완료 후에만 공개되는 구조라고 설명하지 않습니다.',srcs=[GH+'app.py#L1873-L2418'])
for j,(a,b) in enumerate([('CASE','사건·지역 선택'),('EVIDENCE','단서 조사·수집'),('AI','객체 관찰'),('REASON','가설·추리 작성'),('DATA','실제 통계 해석'),('SAFETY','예방 행동 확인')]):
 x=.7+j*2.45;box(s,x,3.35,2.2,2.7,'mint' if j%2==0 else 'ice',radius=True);txt(s,x+.2,3.65,1.9,.4,f'{j+1:02}',17,'teal',True);txt(s,x+.2,4.33,1.9,.45,a,20,'ink',True);txt(s,x+.2,5.12,1.9,.6,b,15,'muted')
 if j<5:txt(s,x+2.23,4.42,.25,.5,'→',18,'teal',True)
bottom(s,'완료 후 실제 통계와 예방 가이드로 연결 · 홈·수사 중에도 데이터 탐색 가능')

s=base('02  CASE DESIGN','네 가지 사건을 실제 통계 유형과 연결합니다','가상 시나리오는 흥미를 만들고, 통계 매핑은 학습 맥락을 만듭니다.',note='app.py CASE_INFO에 정의된 4개 사건입니다. 야간주거침입절도 사건은 지역 통계의 절도 항목과 연결됩니다. 가상 사건의 발생 시각·목격자 진술은 실제 개별 사건 기록이 아닙니다.',srcs=[GH+'app.py#L513-L564'])
for j,(k,t,b,f) in enumerate([('CASE-BURGLARY','심야의 침입자','야간주거침입절도 → 지역 통계: 절도','mint'),('CASE-ROBBERY','사라진 가방','강도 → 지역 통계: 강도','ice'),('CASE-ARSON','불길 속의 흔적','방화 → 지역 통계: 방화','peach'),('CASE-FRAUD','의문의 거래','사기 → 지역 통계: 사기','white')]):
 x=.7+(j%2)*7.45;y=2.65+(j//2)*2.45;card(s,x,y,7.15,2.2,k,t,b,f)
bottom(s,'가상 사건의 단서와 실제 공공 집계 통계는 자료의 성격이 다릅니다.')

s=base('02  LIVE APPLICATION','실제 앱은 사건 선택부터 시작합니다','로컬 실행 화면 · 코드 기준 버전 '+SHA[:7],note='현재 앱을 로컬에서 실행해 캡처했습니다. 앱은 기존 다크 테마이고 이번 산출물은 발표 자료의 밝은 디자인입니다. 원본 앱 UI를 수정한 화면이 아닙니다.',srcs=[GH+'app.py', 'assets/app_home.png'])
addpic(s,AS/'app_home.png',.75,2.6,9.8,5.35)
card(s,10.85,2.75,4.4,4.5,'REAL CAPTURE','진입 장벽을 낮추기','4개 사건 카드\n사건별 설명과 시작 버튼\n홈에서 데이터 탐색','mint')

s=base('02  ARCHITECTURE','사전 가공과 화면 탐색의 책임을 나눕니다','CSV 파이프라인과 이미지 객체 탐지는 서로 다른 입력을 사용합니다.',note='src는 원본 CSV를 processed 산출물로 가공하는 오프라인 스크립트 집합입니다. app.py는 가공 CSV와 사전 학습 YOLO 모델을 읽습니다. 범죄 CSV로 YOLO를 학습하거나 범죄 발생을 예측하는 구조가 아닙니다.',srcs=[GH+'src',GH+'app.py'])
for y,label,items,fill in [(3,'DATA PIPELINE',['공공 통계 CSV','정제·품질 플래그','processed CSV','Streamlit 차트'],'mint'),(5.45,'VISUAL EVIDENCE',['샘플 사건 이미지','yolo11n.pt','객체·신뢰도','증거 후보 수집'],'ice')]:
 txt(s,.8,y-.42,13,.3,label,12,'teal',True)
 for j,t in enumerate(items):
  x=.75+j*3.7;box(s,x,y,3.4,1.35,fill,radius=True);txt(s,x+.2,y+.43,3, .6,t,20,'ink',True)
  if j<3:txt(s,x+3.42,y+.45,.3,.4,'→',18,'teal')
bottom(s,'CSV 분석과 YOLO 관찰을, session_state 기반 게임 흐름에서 연결합니다.')

s=base('03  DATA INVENTORY','파일명보다 실제 데이터 범위를 확인합니다','보관 파일을 직접 읽어 행 수와 기간을 다시 집계했습니다.',note='숫자는 현재 저장된 CSV 실측입니다. 최종 지역 데이터는 56,616행·2011–2024년입니다. 파일명은 1994–2024이지만 1994–2010 행은 들어 있지 않습니다. 지역 75개는 보관된 지역 라벨 수이며 전국 전수나 모든 연도의 완전성을 뜻하지 않습니다. 장기 데이터는 region=범위확인필요이므로 전국 통계로 표기하지 않았습니다.',srcs=['data_audit.json',GH+'processed'])
headers=['보관 데이터','행 수','실제 기간','앱·분석에서의 역할'];widths=[5.2,2,2.9,4.5];xx=[.75,5.95,7.95,10.85]
for x,w,t in zip(xx,widths,headers):box(s,x,2.7,w,.62,'teal');txt(s,x+.15,2.86,w-.3,.35,t,15,'white',True)
rows=[('발생·검거 표준화','9,908','1990–2022','정제 파이프라인 입력'),('REAL DATA Evidence','9,703','1990–2022','증감률·경고·카드 생성'),('CSI Time Machine','9,603','1990–2022','앱 장기 시계열 입력'),('지역 최종 데이터','56,616','2011–2024','앱 지역·연도 탐색')]
for j,row in enumerate(rows):
 for x,w,t in zip(xx,widths,row):box(s,x,3.34+j*.84,w,.8,'white' if j%2==0 else 'mint');txt(s,x+.15,3.55+j*.84,w-.3,.48,t,17,'ink')
bottom(s,'지역 데이터: 7개 도시 · 75개 지역 라벨 · 54개 범죄 유형 · 결측은 별도 확인')

s=base('03  METRIC LITERACY','발생건수·범죄율·검거율·증감률을 구분합니다','범죄율을 이해하는 교육의 출발점은 “무엇을 나누었는가”입니다.',note='범죄율의 교육용 예시는 발생건수/인구×100,000입니다. 현재 앱에는 인구 데이터 결합과 범죄율 계산이 없습니다. 검거율은 공표 값과 계산 값을 구분하며 통계 정의와 집계 시점을 확인해야 합니다. 증감률은 발생건수의 시간 변화이지 개인의 피해 확률이 아닙니다.',srcs=[GH+'src/07_build_analysis_dataset.py',GH+'src/09_build_real_data_evidence.py',GH+'app.py'])
items=[('COUNT · 구현','발생건수','해당 조건에 기록된 절대 건수','1,105건','mint'),('RATE · 교육 개념','인구 보정 범죄율','발생건수 ÷ 인구 × 100,000','현재 앱 미계산','ice'),('CLEARANCE · 구현','검거율','공표 값과 검거/발생 계산 값 구분','자료 정의 확인','peach'),('CHANGE · 구현','발생건수 증감률','(올해−전년) ÷ 전년 × 100','시간 변화 읽기','white')]
for j,(k,t,b,v,f) in enumerate(items):
 x=.7+(j%2)*7.45;y=2.65+(j//2)*2.45;box(s,x,y,7.15,2.22,f,radius=True);txt(s,x+.25,y+.18,6.65,.25,k,11,'teal',True);txt(s,x+.25,y+.66,3.6,.5,t,23,'ink',True);txt(s,x+.25,y+1.4,6.6,.6,b,16,'muted');txt(s,x+3.8,y+.72,3.1,.55,v,19,'blue',True)
bottom(s,'교육 목표는 범죄율 이해 · 현재 앱이 직접 보여주는 지역 지표는 발생건수')

s=base('03  RATE EXAMPLE','같은 비교도 분모를 넣으면 다르게 보입니다','이해를 돕는 가상 계산 예시 · 실제 지역 자료나 현재 앱의 계산 결과가 아닙니다.',note='범죄율 교육을 위한 가상 예시입니다. A는 발생 1,000건·인구 100만 명이므로 인구 10만 명당 100건입니다. B는 발생 200건·인구 10만 명이므로 10만 명당 200건입니다. 절대 건수와 인구 보정 비율의 비교 방향이 달라질 수 있음을 보여줍니다. 이 수치가 실제 지역이나 위험 순위를 뜻하지 않으며, 신고·집계 정의·유동인구 등도 함께 확인해야 합니다. 현재 앱에는 인구 결합 범죄율 계산이 없습니다.',srcs=[GH+'app.py#L2370-L2380'])
for x,label,count,pop,rate,fill in [(.75,'가상 A 지역','1,000건','100만 명','100건 / 10만 명','mint'),(8.15,'가상 B 지역','200건','10만 명','200건 / 10만 명','ice')]:
 box(s,x,2.75,7.05,3.95,fill,radius=True);txt(s,x+.3,3.02,6.5,.4,label,21,'teal',True);txt(s,x+.3,3.86,3,.75,count,34,'ink',True);txt(s,x+3.7,4.0,3,.5,'인구 '+pop,21,'muted');line(s,x+.3,4.85,x+6.65,4.85);txt(s,x+.3,5.24,6.4,.8,rate,30,'blue',True)
bottom(s,'교육용 산식: 발생건수 ÷ 인구 × 100,000 · 비교 기간·집계 기준도 같아야 합니다.')

s=base('03  DATA PIPELINE','원본을 바꾸기보다, 비교 가능한 정보를 추가합니다','표준화 → 세부 범죄 분리 → 변화율 → 커버리지 → 게임 연결',note='9,908행에서 세부 분석 대상 등을 거쳐 Evidence 9,703행이 생성됩니다. 이 중 발생건수 존재 행 9,603개로 Time Machine을 구성하고 CORE 10종×33년=330행을 선택합니다. 행 수 감소를 모두 오류 제거 또는 중복 제거라고 표현하지 않습니다.',srcs=['data_audit.json',GH+'src/09_build_real_data_evidence.py',GH+'src/10_build_time_machine.py',GH+'src/11_build_csi_time_machine.py'])
for j,(n,t,b) in enumerate([('9,908','표준화','원천별 공통 컬럼'),('9,703','Evidence','세부 분석·통계 존재'),('9,603','Time Machine','발생건수 존재'),('330','CORE','10개 죄종 × 33년')]):
 x=.75+j*3.7;box(s,x,3,3.4,3.75,'white',radius=True);txt(s,x+.25,3.4,2.9,.85,n,37,'teal',True);txt(s,x+.25,4.6,2.9,.5,t,22,'ink',True);txt(s,x+.25,5.55,2.9,.7,b,17,'muted')
bottom(s,'기간의 길이뿐 아니라 결측·연속성·분류 체계도 함께 확인합니다.')

code('04  CORE CODE · DATA','결측 기호를 0으로 바꾸지 않습니다','숫자 파싱의 첫 단계에서 “자료 없음”의 의미를 보존합니다.','src/05_clean_occurrence_arrest.py',39,55,'“-”, 빈칸, NaN은 결측으로 보존','쉼표가 있는 숫자는 float로 변환','관측되지 않은 값과 실제 0건을 구분')
code('04  CORE CODE · DATA','연도별 넓은 표를 관측값 단위로 바꿉니다','Wide → Long 변환으로 시간 축 필터·그룹 연산의 기반을 만듭니다.','src/05_clean_occurrence_arrest.py',123,129,'melt로 연도 컬럼을 행으로 변환','죄종·지표·연도별 관측값 구성','후속 pivot으로 공통 지표 컬럼 생성',note='실제 melt 발췌입니다. 이후 연도 추출과 숫자 정리를 하고 year, crime_type 키로 pivot_table(aggfunc=first)을 수행합니다. first 집계는 중복이 없음을 자동으로 보증하지 않으므로 품질 검사가 별도로 필요합니다.')
code('04  CORE CODE · DATA','상위 합계와 세부 범죄를 분리합니다','집계행을 개별 범죄 합산에 다시 넣는 문제를 줄입니다.','src/08_2_refine_crime_groups.py',335,337,'is_aggregate의 반대로 세부 분석 대상 표시','Evidence 생성 때 이 플래그로 필터','범죄군 재분류는 규칙 기반이며 공식 분류와 구분',note='AGGREGATE_NAMES 집합으로 상위 집계명을 판별하고 include_detail_analysis=~is_aggregate를 만듭니다. Evidence 생성 스크립트에서 True 행만 선택합니다. 이후 범죄명 키워드 분류는 프로젝트용 보조 범죄군이며 공공 통계의 원래 분류를 대체하는 공식 분류가 아닙니다.')
code('04  CORE CODE · QUALITY','공표 검거율과 계산 검거율을 구분합니다','원본 clearance_rate는 유지하고 비교용 값을 추가합니다.','src/07_build_analysis_dataset.py',64,74,'발생건수 > 0일 때 검거/발생×100','원본과 계산 값의 차이를 확인 가능','100% 초과는 자동 삭제보다 정의·집계 기준 확인',note='계산 검거율은 발생과 검거가 있고 발생건수가 양수일 때 생성합니다. 원본 공표 검거율은 덮어쓰지 않습니다. 검거율 100% 초과 또는 불일치는 품질 플래그로 남깁니다. 연간 발생과 검거가 같은 사건 집합인지 등 원자료 정의 확인이 필요합니다.')
code('04  CORE CODE · TIME','전년이 실제로 존재하는지 확인합니다','직전 관측 행이 항상 직전 연도를 의미하지는 않습니다.','src/09_build_real_data_evidence.py',101,115,'이전 연도 = 올해−1 조건 검사','연도 공백이 있으면 전년 비교를 생략','이전 값 양수·현재 값 존재 조건도 함께 적용')
code('04  CORE CODE · TIME','비교 가능한 경우에만 증감률을 계산합니다','분모와 연속성이 유효한 행에만 전년 대비 변화율을 기록합니다.','src/09_build_real_data_evidence.py',118,130,'np.where(valid_change, 계산식, NaN)','결측·연도 공백을 억지 수치로 표시하지 않음','증감률은 발생건수의 변화이며 범죄율과 다름')
code('04  CORE CODE · COVERAGE','긴 시계열에도 관측 완성도 기준을 둡니다','연도 수와 커버리지를 함께 보아 분석 등급을 나눕니다.','src/10_build_time_machine.py',99,113,'30Y: 30년 이상·커버리지 90% 이상','10개 CORE 죄종은 각 33년 관측','등급은 기간 적합성 기준이며 통계 품질 보증은 아님',note='coverage_rate=year_count/year_span×100입니다. 30Y 30년·90%, 20Y 20년·80%, 10Y 10년·70% 기준을 사용합니다. 33년을 관측했어도 법·집계 기준이 기간 내내 동일하다고 보장하지는 않습니다.')

s=base('04  DATA REVIEW','데이터를 신뢰하려면, 남은 문제도 밝혀야 합니다','현재 산출물에서 확인한 범위와 코드상 개선 과제를 분리했습니다.',note='현행 지역 CSV를 직접 확인하면 기간은 2011–2024뿐입니다. src/20의 파일 순회 뒤 후속 처리가 루프 바깥에 있어 마지막 파일만 결과에 반영되는 구조를 확인했습니다. make_district_name은 군으로 끝나지 않는 모든 이름에 구를 붙이므로 중구구 등 중복 표기를 만듭니다. 이번 작업은 발표 문서 개편이므로 앱과 데이터는 수정하지 않았습니다.',srcs=['data_audit.json',GH+'src/20_build_final_region_dataset.py#L205-L315',GH+'src/20_build_final_region_dataset.py#L471-L478'])
for j,(k,t,b,f) in enumerate([('01 · PERIOD','지역 실제 기간 2011–2024','파일명 1994–2024와 다름\n루프 밖 후속 처리 구조 확인','mint'),('02 · LABEL','일부 지역명 접미사 중복','중구구·동구구 등 보관\n지역 표시 함수 개선 필요','ice'),('03 · SCOPE','장기 통계 집계 범위 확인','region = 범위확인필요\n전국 범죄율로 표기하지 않음','peach')]):card(s,.7+j*4.95,2.8,4.7,4.25,k,t,b,f)
bottom(s,'자료의 범위를 정확히 말하는 것 자체가 데이터 리터러시의 일부입니다.')

code('04  CORE CODE · REGION','같은 지역명도 도시와 함께 복원합니다','원본 지역 목록과 OLD/NEW 구조가 일치하는지 먼저 비교합니다.','src/20_build_final_region_dataset.py',292,300,'실제 지역 순서를 사전 정의한 목록과 비교','일치하는 구조를 골라 도시 이름 연결','불일치 시 city 결측·region_match=False로 보류',note='중구·강서처럼 여러 도시에 있는 이름을 이름 하나만으로 매핑하지 않습니다. 원본 순서와 전체 목록을 비교합니다. 다만 현재 저장된 region_match=True는 정의한 목록과 일치했다는 뜻이며 외부 행정구역 원문 대조를 대체하지 않습니다.')
code('04  CORE CODE · REGION','관측값의 키를 네 가지 조건으로 정의합니다','연도 + 도시 + 구·군 + 범죄 유형이 하나의 관측 단위입니다.','src/20_build_final_region_dataset.py',498,511,'최종 키 중복을 keep=False로 표시','현재 저장 CSV의 동일 키 중복 0행','중복 0과 결측 0·표기 정확성은 별도 문제',note='최종 CSV를 직접 읽어 같은 키로 재검사했을 때 중복 행은 0입니다. 반면 발생건수 결측은 15,888행입니다. 중복이 없다는 사실을 데이터가 완전하거나 오류가 없다는 주장으로 확대하지 않습니다.')
code('05  CORE CODE · INTERACTION','화면 이동과 사건 상태를 함께 관리합니다','Streamlit 재실행 사이의 진행 정보를 session_state에 둡니다.','app.py',666,672,'사건 코드와 페이지를 session_state에 저장','새 사건을 열 때 수집 증거 목록 초기화','인증·DB 저장이 없는 세션 단위 학습 프로토타입',note='open_case 함수의 실제 발췌입니다. selected_case, page, investigation_started, investigation_view, collected_evidence를 설정합니다. 일부 다른 상태는 이 함수에서 초기화되지 않으므로 여러 사건을 연속 수행할 때 전체 상태 초기화는 추가 검증 과제입니다.')

s=base('05  EVIDENCE LAB','증거를 모으는 과정이 학습의 입력이 됩니다','실제 실행 화면 · CCTV·현장 물체·시간·목격자·지역·과거 기록',note='화면의 조사 경로는 CCTV, 현장 물체, 시간, 목격자, 지역 데이터, 과거 기록의 여섯 가지입니다. 일부 안내 문구에는 YOLO 연결 예정이라고 남아 있지만 실제 버튼 경로에는 YOLO 분석이 구현되어 있으므로 실행된 코드 기준으로 설명했습니다.',srcs=[GH+'app.py#L1060-L1187','assets/app_investigation.png'])
addpic(s,AS/'app_investigation.png',.75,2.6,9.8,5.4)
card(s,10.85,2.75,4.4,4.6,'EXPLORE','여섯 가지 조사 경로','이미지에서 관찰하기\n진술과 비교하기\n지역·과거 통계 확인하기','ice')

code('05  CORE CODE · LEARNING','가설 전에 최소한의 관찰을 요구합니다','증거 3개 이상을 모은 뒤 가설 작성 UI를 엽니다.','app.py',2063,2072,'수집 증거 수가 3개 미만이면 안내','단서를 모으기 전에 결론부터 쓰지 않도록 유도','3개 조건은 학습 흐름이며 증거의 충분성 판정은 아님')
code('05  CORE CODE · AI','모델은 한 번 읽고 재사용합니다','cache_resource로 YOLO 모델 객체를 유지합니다.','app.py',10,14,'번들 yolo11n.pt를 YOLO로 로드','화면 재실행 시 모델 재로딩 부담 감소','로컬 모델 사용 · 범죄 CSV로 재학습하지 않음')
code('05  CORE CODE · AI','CCTV 프레임에서 객체를 관찰합니다','모델 출력은 물체의 종류·박스·신뢰도입니다.','app.py',1359,1369,'yolo_model에 CCTV 이미지 경로 전달','박스를 그린 뒤 BGR→RGB로 화면에 표시','탐지 결과는 시각 단서 후보이며 범죄 판단이 아님')

s=base('05  AI EVIDENCE','AI는 “무엇이 보이는가”를 돕습니다','이 프레임에서 관찰한 객체를 사건 관련성과 구분합니다.',note='번들 yolo11n.pt와 img/cctv_burglary.png로 CPU 추론한 결과입니다. 5개 객체 박스가 반환되었습니다. 이는 단일 샘플의 기능 확인이며 정확도·재현율·사건 판단 성능의 평가가 아닙니다. 표시 신뢰도는 물체 클래스에 대한 모델 점수입니다.',srcs=[GH+'yolo11n.pt',GH+'img/cctv_burglary.png','validation.json'])
addpic(s,AS/'yolo_annotated.jpg',.8,2.65,9.7,4.9)
card(s,10.8,2.8,4.45,4.3,'LOCAL CPU INFERENCE','객체 5개 관찰','클래스·박스·신뢰도\n관찰된 물체를 증거 후보로\n사람의 역할·범행은 판단하지 않음','mint')
bottom(s,'신뢰도 90% = 범죄 확률 90%가 아닙니다.','peach','orange')

code('05  CORE CODE · AI OUTPUT','탐지 클래스와 신뢰도를 설명 가능한 값으로 바꿉니다','모델의 숫자 출력을 클래스 이름·점수로 읽습니다.','app.py',1376,1387,'box.cls와 box.conf를 숫자로 추출','result.names에서 객체 이름 조회','한국어 매핑과 함께 사용자에게 관찰 정보 전달')
code('05  CORE CODE · EVIDENCE','같은 증거 후보를 반복해서 쌓지 않습니다','수집 버튼은 목록 포함 여부를 확인한 뒤 append합니다.','app.py',1485,1496,'증거 이름이 없을 때만 목록 추가','같은 이름의 중복 클릭으로 점수 누적 방지','문자열 단위 중복 확인 · 고유 증거 ID는 개선 과제')
code('05  CORE CODE · COMPARISON','객체와 진술을 간단한 규칙으로 비교합니다','가방 관련 문자열이 양쪽에 있는지 검사합니다.','app.py',1971,1978,'CCTV 증거 안의 가방 키워드를 확인','목격자 증거도 같은 방식으로 비교','단순 키워드 비교이며 인물 동일성 추론은 아님',note='예시는 cctv_has_bag 계산 부분입니다. 이어서 witness_has_bag를 같은 방식으로 계산하고 둘 다 True일 때 공통 관찰 정보로 가방을 추가합니다. 외부 LLM이나 인물 재식별 모델을 사용한 비교가 아닙니다.')
code('05  CORE CODE · FEEDBACK','점수는 학습 활동의 진행도를 보여줍니다','증거 최대 50점 + 가설 20점 + 최종 추리 30점','app.py',2213,2228,'증거 수와 작성 여부를 점수로 합산','추가 관찰과 설명 작성을 유도','정답·사실성·글의 품질을 자동 채점하는 점수는 아님')
code('06  CORE CODE · REAL DATA','완료한 사건의 조건으로 실제 통계를 찾습니다','지역·구·군·범죄 유형을 같은 조건으로 연결합니다.','app.py',2292,2299,'도시·구군·region_crime의 동시 필터','가상 CASE와 관련된 지역 집계 시계열 표시','실제 개별 사건과 일대일 매칭한 기록은 아님')

local=pd.read_csv(ROOT/'processed/region_crime_final_1994_2024.csv');q=local[(local.city=='서울특별시')&(local.district=='용산구')&(local.crime_type_original=='절도')].sort_values('year')
s=base('06  READ THE PATTERN','짧게는 증가, 길게는 감소일 수 있습니다','서울특별시 용산구 · 절도 발생건수 · 2011–2024',note='현재 저장 CSV의 실제 값으로 차트를 만들었습니다. 2011년 1,603건, 2023년 987건, 2024년 1,105건입니다. 2024년 전년 대비 약 +12.0%, 2011년 대비 약 −31.1%입니다. 이 계산은 발표용 재계산이며 앱에 새 계산 기능을 추가한 것은 아닙니다. 원자료 정의와 분류·신고·집계 변화, 인구를 확인하지 않고 위험 변화나 원인을 단정하지 않습니다.',srcs=[GH+'processed/region_crime_final_1994_2024.csv','data_audit.json'])
box(s,.75,2.7,9.65,4.6,'white',radius=True);chart(s,.95,2.9,9.15,4.15,q.year.tolist(),[('절도',q.occurrence_count.tolist())])
for y,num,label,col in [(2.9,'1,105건','2024년 발생건수','teal'),(4.35,'+12.0%','2023년 대비 · 발표 재계산','blue'),(5.8,'−31.1%','2011년 대비 · 발표 재계산','orange')]:txt(s,10.9,y,4.2,.7,num,35,col,True);txt(s,10.92,y+.8,4.2,.45,label,15,'muted')
bottom(s,'어떤 기간을 비교했는지 먼저 말하고, 같은 지표로 흐름을 설명합니다.')

s=base('06  DATA LITERACY ACTIVITY','차트를 읽고, 근거가 있는 한 문장을 씁니다','수업에서 사용할 수 있는 활동 제안 · 앱의 자동 채점 기능은 아닙니다.',note='교육 활동으로 제안하는 슬라이드입니다. 1. 2024년의 건수와 단위를 말한다. 2. 전년과 시작연도를 각각 비교한다. 3. 인구 보정이 없으므로 이 숫자만으로 위험 증가를 말할 수 없다고 설명한다. 4. 변화 원인을 알려면 신고·집계 기준·인구와 노출 규모 등을 추가 확인한다. 이 질문들은 신규 앱 기능으로 구현한 것이 아닙니다.',srcs=['data_audit.json',GH+'app.py#L2370-L2380'])
for j,(k,t,b,f) in enumerate([('STEP 01','무엇이 달라졌나?','1,603건 → 1,105건\n비교 기간: 2011–2024','mint'),('STEP 02','비교 창을 바꾸면?','2023년 987건 → 2024년 1,105건\n같은 자료도 단기·장기 방향이 다름','ice'),('STEP 03','무엇을 더 알아야 할까?','인구·노출 규모·분류·집계 기준\n원인과 위험은 추가 근거가 필요','peach')]):card(s,.7+j*4.95,2.8,4.7,4.25,k,t,b,f)
bottom(s,'예시 답변: “최근에는 늘었지만 2011년보다는 적다. 위험 변화는 추가 자료가 필요하다.”')

code('06  CORE CODE · BASELINE','기준연도 대비 변화량도 따로 계산합니다','장기 비교는 무엇을 기준으로 했는지 함께 제시합니다.','src/11_build_csi_time_machine.py',162,169,'1990 기준값으로 발생건수 변화율 산출','전년 변화와 장기 변화를 나눠 읽기','기준값 0 처리·집계 범위 검증은 추가 개선 필요',note='해당 코드의 change_from_1990_pct 산식입니다. 1990 기준값은 CORE 행에서 추출해 죄종명으로 merge합니다. 현재 산식에는 기준값 0에 대한 명시적 마스킹이 없습니다. 장기 CSV의 집계 지역은 범위확인필요이므로 전국 변화로 해석하지 않습니다.')
code('06  CORE CODE · LOCAL EXPLORER','사용자 선택을 같은 조건의 시계열로 연결합니다','도시 → 구·군 → 범죄 → 연도 순서로 탐색합니다.','app.py',2925,2936,'선택한 세 조건으로 지역 데이터를 필터','연도 정렬 후 슬라이더·선 차트에 연결','관측값 결측은 0건으로 해석하지 않음')

s=base('06  SAFETY LAB','데이터 해석을 일상 행동으로 연결합니다','범죄 유형별 일반 예방 가이드를 확인합니다.',note='SAFETY_GUIDE는 코드에 작성된 범죄 유형별 일반 안내입니다. 통계에서 학습한 개인별 위험 예측이나 지역 맞춤 처방이 아닙니다. 예방 가이드가 데이터 기반 추천 모델의 출력이라고 표현하지 않습니다.',srcs=[GH+'app.py#L599-L636',GH+'app.py#L2394-L2418'])
for j,(k,t,b,f) in enumerate([('THEFT','잠금과 보관','출입문·창문 잠금 확인\n귀중품이 밖에서 보이지 않게','mint'),('FRAUD','연락과 거래 확인','발신자를 별도로 재확인\n송금 전 거래 정보 확인','ice'),('FIRE','대피와 신고','비상구·대피 경로 확인\n안전한 장소에서 119 신고','peach')]):card(s,.7+j*4.95,2.8,4.7,4.25,k,t,b,f)
bottom(s,'지역 위험 등급을 매기기보다, 일반적인 예방 행동을 함께 살펴봅니다.')

s=base('07  ENGINEERING EVIDENCE','기능은 실제 실행 경로로 확인했습니다','로컬 CPU 환경 · 현재 코드와 번들 데이터 기준',note='Streamlit AppTest에서 실제 버튼과 텍스트 입력을 사용해 홈→사건 선택→시작→CCTV→YOLO→증거 수집→목격자 수집→가설 저장→최종 추리 제출→REAL DATA UNLOCKED와 Safety LAB 렌더링을 확인했습니다. 대표 burglary 1개 경로의 검증이며 모든 사건과 모든 상태 조합을 검증한 것은 아닙니다. 별도 CPU 추론에서 5개 객체 박스를 확인했고 pip check 및 HTTP health/root가 성공했습니다.',srcs=['validation.json','data_audit.json',GH+'app.py'])
rows=[('PASS','앱 실행','HTTP health / 루트 응답 확인'),('PASS','사용자 경로','실제 버튼·입력으로 추리 제출까지'),('PASS','통계·가이드','REAL DATA UNLOCKED / SAFETY LAB 표시'),('PASS','이미지 분석','번들 모델 CPU 추론 · 객체 5개'),('REVIEW','데이터 품질','기간·지역명·집계 범위 개선 과제 확인')]
for j,(status,t,b) in enumerate(rows):
 y=2.65+j*.9;box(s,.75,y,14.5,.76,'white',radius=True);pill(s,.95,y+.18,1.3,status,'mint' if status=='PASS' else 'peach','teal' if status=='PASS' else 'orange');txt(s,2.65,y+.17,3.3,.4,t,18,'ink',True);txt(s,6.2,y+.2,8.75,.4,b,16,'muted')
bottom(s,'기능 검증과 효과 검증은 구분: 탐지 정확도·학습 향상·전 사건 완주는 미측정','peach','orange')

s=base('07  IMPLEMENTED / NEXT','구현한 기능과 교육 효과의 입증을 구분합니다','현재의 강점을 유지하면서, 다음 검증의 순서를 정합니다.',note='구현 완료는 코드 또는 실행으로 확인된 것만 표기했습니다. 교육 효과 향상, 피해 감소율, AI 정확도 등은 근거가 없어 수치로 주장하지 않았습니다. 현재 앱의 한계와 개선 계획은 향후 작업이며 이번 발표 개편에서 앱을 고친 것은 아닙니다.',srcs=[GH+'app.py','validation.json','data_audit.json'])
card(s,.75,2.8,7.1,4.3,'IMPLEMENTED','몰입 → 관찰 → 데이터','4개 CASE와 증거 수집\nYOLO 객체 관찰·가설·추리\n지역·장기 차트·예방 가이드','mint')
card(s,8.1,2.8,7.1,4.3,'NEXT VALIDATION','품질 → 이해 → 평가','자료 범위·지역명 교정\n인구 보정 범죄율 결합\n사전·사후 문항으로 이해도 측정','ice')
bottom(s,'“재미있다”에서 “통계를 근거로 설명할 수 있다”로 학습 목표를 평가합니다.')

s=base('07  ROADMAP','다음 개선은 더 정확한 근거에서 시작합니다','발표를 위한 계획 · 아직 구현하거나 성과를 측정한 항목이 아닙니다.',note='1순위는 기존 데이터 품질과 출처 메타데이터의 정리입니다. 2순위는 동일 기간·행정구역 인구를 결합한 범죄율과 집계 단위 안내입니다. 3순위는 교육 활동을 문항과 루브릭으로 검증하는 것입니다. UI 디자인 개선과 세션 초기화 검증도 운영 품질 과제로 남깁니다.',srcs=['data_audit.json',GH+'src/20_build_final_region_dataset.py',GH+'app.py'])
for j,(k,t,b,f) in enumerate([('01 · DATA QUALITY','자료 기준부터 정리','원문 URL·지역·단위 기록\n지역 생성 루프·접미사 개선\n결측·분류 변경 이력 표시','mint'),('02 · METRIC LITERACY','범죄율 비교 확장','같은 기간·지역 인구 결합\n건수와 비율을 나란히 표시\n노출 규모·해석 한계 설명','ice'),('03 · LEARNING VALIDATION','이해도를 확인','사전·사후 문항 설계\n근거 설명 루브릭 도입\n사건 간 상태 초기화 검증','peach')]):card(s,.7+j*4.95,2.8,4.7,4.25,k,t,b,f)

s=base('PROJECT MESSAGE','사건을 풀다가, 데이터를 읽는 사람이 됩니다','범죄 데이터를 쉽게 이해하고, 패턴을 근거로 설명하는 교육 경험',note='CSI LAB의 핵심 메시지입니다. 현재 구현은 발생건수 탐색과 객체 관찰·가설 작성의 프로토타입입니다. 범죄율 이해와 패턴 읽기가 교육 목표이며 위험 예측 성능이나 교육 효과를 검증한 서비스라고 주장하지 않습니다.',srcs=[GH+'README.md',GH+'app.py'])
txt(s,.8,3.15,13.8,1.6,'관찰하고, 비교하고,\n근거를 설명합니다.',43,'teal',True)
for x,a,b in [(.85,'OBSERVE','무엇이 보이는가'),(5.75,'READ','어떻게 달라졌는가'),(10.65,'EXPLAIN','무엇을 말할 수 있는가')]:
 pill(s,x,6.05,4.45,a,'mint' if x<5 else 'ice');txt(s,x,6.65,4.5,.7,b,22,'ink',True)

s=base('APPENDIX  /  SOURCE & METHOD','모든 주장을 코드와 보관 데이터로 되돌릴 수 있습니다','참고자료는 구성만 참고 · 내용·수치·코드는 CSI LAB 저장소에서 확인',note='참고 파일 where_u_go.pptx(37장)는 구성, 코드 중심 설명, 가독성 수준만 참고했습니다. 타 프로젝트의 데이터·모델·성과·기여자 이름은 CSI LAB에 옮기지 않았습니다. 모든 코드 발췌는 원본 행 번호를 보존한 링크와 발표 노트에 수록했습니다. 자료 제작 기준 커밋은 '+SHA+'입니다. 산출물의 데이터 수치는 data_audit.json으로 재확인할 수 있습니다. 개별 공공 통계의 원문 URL·추출일·집계 범위는 보관 파일만으로 확정하지 않았습니다.',srcs=[GH+'app.py',GH+'src',GH+'data/raw/crime',GH+'processed','data_audit.json','validation.json'])
for j,(k,t,b,f) in enumerate([('CODE','실제 구현 근거','app.py · src/*.py\n발췌 파일·행 번호·원본 링크\n전체 설명은 발표 노트','mint'),('DATA','직접 다시 계산한 근거','보관 CSV 행 수·기간·결측\n용산구 절도 차트·변화율\nJSON에 파일 SHA-256 기록','ice'),('REFERENCE','첨부자료의 역할','where_u_go.pptx: 구성 참고\n프로젝트 내용·수치는 미전용\n실제 캡처는 현재 로컬 실행','peach')]):card(s,.7+j*4.95,2.8,4.7,4.25,k,t,b,f)
bottom(s,'소스 기준 '+SHA[:7]+' · 교육 효과·탐지 정확도·범죄율 산출은 미검증 또는 미구현')

# Write editable deck and companion evidence.
# Disable theme shadows for a clean, bright document style.
for slide in prs.slides:
 for shape in slide.shapes:
  for effect in shape._element.xpath('.//a:effectRef'):
   effect.set('idx', '0')
path=OUT/'CSI_LAB_교육프로젝트_전면개편.pptx';prs.save(path)
(OUT/'code_excerpts.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2)+'\n')
md=['# CSI LAB — 발표 원고와 근거','', '범죄율을 올바르게 이해하고 범죄 데이터의 패턴을 읽는 교육 프로젝트.', '',f'구현 기준: `{SHA}`. 첨부 PPT는 구성과 설명 수준만 참고했습니다.','', '현재 지역 지표는 발생건수이며 인구 보정 범죄율은 미구현입니다.','']
for n in narrative:
 md+=['## '+str(n['slide']).zfill(2)+'. '+n['title'],'',n['subtitle'],'',n['notes'],'']
 for e in sources:
  if e['slide']==n['slide']:md+=['```python',e['excerpt'],'```','',f"원본: [{e['file']} L{e['start']}–L{e['end']}]({e['url']})",'']
 if n['sources']:md+=['근거: '+' · '.join(n['sources']),'']
(OUT/'발표원고_근거.md').write_text('\n'.join(md)+'\n')
print(f'Saved {path}: {len(prs.slides)} slides, {len(sources)} source excerpts')
